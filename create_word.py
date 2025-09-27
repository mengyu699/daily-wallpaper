#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
import re

# 创建Word文档
doc = Document()

# 设置文档标题
title = doc.add_heading('垂直小号-公众号丨实战手册丨2025年9月航海', level=0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# 读取markdown内容
with open('/Users/mengyu/Desktop/Cursor/course_content.md', 'r', encoding='utf-8') as f:
    content = f.read()

# 按行处理内容
lines = content.split('\n')
i = 0
while i < len(lines):
    line = lines[i].strip()
    
    if not line:
        i += 1
        continue
    
    # 处理标题
    if line.startswith('###'):
        if line.startswith('#### '):
            text = line.replace('#### ', '')
            doc.add_heading(text, level=4)
        elif line.startswith('### '):
            text = line.replace('### ', '')
            doc.add_heading(text, level=3)
    elif line.startswith('##'):
        if line.startswith('##### '):
            text = line.replace('##### ', '')
            doc.add_heading(text, level=5)
        elif line.startswith('## '):
            text = line.replace('## ', '')
            doc.add_heading(text, level=2)
    elif line.startswith('#'):
        if not line.startswith('# 垂直小号'):  # 跳过重复的主标题
            text = line.replace('# ', '')
            doc.add_heading(text, level=1)
    
    # 处理粗体文本
    elif line.startswith('**') and line.endswith('**'):
        text = line.replace('**', '')
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
    
    # 处理列表项
    elif line.startswith('• '):
        text = line.replace('• ', '')
        p = doc.add_paragraph(text, style='List Bullet')
    elif line.startswith('- '):
        text = line.replace('- ', '')
        p = doc.add_paragraph(text, style='List Bullet')
    elif re.match(r'^\d+\. ', line):
        p = doc.add_paragraph(line, style='List Number')
    
    # 处理分隔线
    elif line == '---':
        doc.add_page_break()
    
    # 处理普通段落
    else:
        # 处理段落中的格式
        p = doc.add_paragraph()
        
        # 分割粗体文本
        parts = re.split(r'(\*\*[^*]+\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                text = part[2:-2]
                run = p.add_run(text)
                run.bold = True
            elif part:
                p.add_run(part)
    
    i += 1

# 保存文档
doc.save('/Users/mengyu/Desktop/Cursor/course_content.docx')
print("Word文档已成功创建：/Users/mengyu/Desktop/Cursor/course_content.docx")